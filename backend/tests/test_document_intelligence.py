"""End-to-end coverage for clause-level document intelligence.

These tests drive the real HTTP surface with the original and revised Website Development
Agreement fixtures, so they cover upload validation, clause extraction, persistent statuses,
version comparison, per-finding Q&A, and the PDF report together.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.providers.base import AIResult
from app.providers.mock import MockAIProvider
from tests.conftest import register

FIXTURES = Path(__file__).parent / "fixtures"
ORIGINAL = FIXTURES / "website_development_agreement_original.txt"
REVISED = FIXTURES / "website_development_agreement_revised.txt"


def collapse(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


class AnsweringProvider(MockAIProvider):
    """A development provider that can also answer questions about a finding."""

    name = "test-answerer"
    model = "test-model"

    answer = "The clause leaves the period open to argument."
    revision = "The Client shall pay each invoice within 30 days."

    def answer_finding_question(self, finding, question, clause_text="", history=None, request_revision=False):
        return AIResult(
            data={
                "answer": f"{type(self).answer} ({question})",
                "suggested_revision": type(self).revision if request_revision else "",
            },
            provider=self.name,
            model=self.model,
        )


class SilentProvider(MockAIProvider):
    name = "test-silent"
    model = "test-model"

    def answer_finding_question(self, finding, question, clause_text="", history=None, request_revision=False):
        return AIResult(data={"answer": "   "}, provider=self.name, model=self.model)


def build_two_page_pdf() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    for line in (
        "PAGE ONE SERVICE AGREEMENT",
        "1. Scope of Services",
        "The Supplier shall deliver the Services within four weeks.",
        "2. Payment Terms",
        "The Customer shall pay each invoice within 14 days.",
    ):
        pdf.drawString(72, 720, line)
    pdf.showPage()
    for line in (
        "PAGE TWO",
        "3. Intellectual Property",
        "All rights in the deliverables vest in the Supplier.",
        "4. Termination",
        "Either party may terminate on 30 days written notice.",
    ):
        pdf.drawString(72, 720, line)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def analyse(client, path: Path, **form) -> dict:
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": (path.name, path.read_bytes(), "text/plain")},
        data=form,
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200, response.text
    return response.json()


def analyse_original(client, **form) -> dict:
    return analyse(client, ORIGINAL, **form)


def analyse_revised(client, document_id: int, previous_review_id: int) -> dict:
    return analyse(
        client,
        REVISED,
        document_id=str(document_id),
        previous_review_id=str(previous_review_id),
    )


def make_document(client, title: str = "Website Development Agreement") -> int:
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/documents/generate",
        json={
            "template_slug": "general-agreement",
            "fields": {
                "document_title": title,
                "party_one": "Brightlane Digital Ltd",
                "party_two": "Harborline Trading Co",
                "purpose": "Website development",
                "obligations_one": "Design and deliver the website",
                "obligations_two": "Provide timely feedback",
                "term": "12 weeks",
                "termination": "30 days notice",
                "governing_law": "England and Wales",
            },
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 201, response.text
    return response.json()["document_id"]


# --------------------------------------------------------------------------- upload validation


def test_review_rejects_disallowed_file_extension(client):
    register(client)
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": ("payload.exe", b"MZ\x90\x00", "application/octet-stream")},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "invalid_upload"


def test_review_rejects_text_declared_as_pdf(client):
    """A .pdf name must not let arbitrary text through the PDF reader."""

    register(client)
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": ("agreement.pdf", b"this is plainly not a pdf", "application/pdf")},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "invalid_upload"


def test_review_rejects_docx_with_wrong_signature(client):
    register(client)
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": ("agreement.docx", b"PK\x03\x04 but not a real docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "invalid_upload"


def test_review_rejects_oversized_upload(client):
    register(client)
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    oversized = b"clause text. " * 2_000_000
    response = client.post(
        "/api/review",
        files={"file": ("big.txt", oversized, "text/plain")},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 400
    assert "MB" in response.json()["detail"]["message"]


def test_review_accepts_a_real_pdf_with_exact_page_references(client):
    register(client)
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": ("agreement.pdf", build_two_page_pdf(), "application/pdf")},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["file_type"] == "pdf"
    assert body["page_count"] == 2
    assert body["page_reference_kind"] == "source_page"
    clause_pages = {clause["page_start"] for clause in body["clauses"]}
    assert clause_pages <= {1, 2}
    assert 2 in clause_pages, "page two content must be cited on page 2"

    intellectual_property = next(
        clause for clause in body["clauses"] if "Intellectual Property" in clause["heading"]
    )
    assert intellectual_property["page_start"] == 2
    assert intellectual_property["citation"] == "3. Intellectual Property (Page 2)"


def test_review_accepts_docx_and_marks_pages_as_estimated(client):
    register(client)
    document = DocxDocument()
    document.add_heading("CONSULTING AGREEMENT", level=1)
    document.add_paragraph("1. Scope")
    document.add_paragraph("The Consultant shall provide the Services for a fixed fee of GBP 20,000.")
    document.add_paragraph("2. Payment")
    document.add_paragraph("The Client shall pay within 30 days of invoice.")
    buffer = io.BytesIO()
    document.save(buffer)

    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={
            "file": (
                "agreement.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["file_type"] == "docx"
    assert body["page_reference_kind"] == "estimated_page"
    assert "Services" in body["extracted_text"]


# --------------------------------------------------------------------------- clause level review


def test_original_agreement_produces_clause_references_and_typed_findings(client):
    register(client)
    body = analyse_original(client)

    assert body["clause_count"] >= 10
    assert body["page_reference_kind"] == "estimated_page"
    assert len(body["clauses"]) == body["clause_count"]
    headings = " ".join(clause["heading"] for clause in body["clauses"])
    assert "Fees and Payment" in headings
    assert "Governing Law" in headings

    for finding in body["detailed_findings"]:
        assert finding["clause_citation"]
        assert finding["clause_text"], "each finding must carry the clause it came from"
        assert finding["page_start"] >= 1
        assert finding["status"] == "open"
        assert finding["category"] in {
            "unclear_clause",
            "missing_information",
            "inconsistent_detail",
            "ambiguous_wording",
            "unusual_obligation",
            "professional_review",
        }
        assert finding["severity"] in {"low", "medium", "high"}
        assert collapse(finding["excerpt"]) in collapse(finding["clause_text"])

    assert "not legal advice" in body["disclaimer"]


def test_findings_point_at_the_specific_sentence_that_triggers_them(client):
    register(client)
    body = analyse_original(client)
    titles = {finding["title"] for finding in body["detailed_findings"]}
    assert any("90 days" in title for title in titles)
    assert any("non-refundable" in title for title in titles)

    payment = next(
        finding
        for finding in body["detailed_findings"]
        if "90 days" in finding["title"]
    )
    assert "within 90 days" in payment["excerpt"]
    assert payment["page_start"] >= 1


def test_quote_offsets_are_relative_to_the_stored_clause_text(client):
    """A consumer must be able to highlight the excerpt straight out of clause_text."""

    register(client)
    body = analyse_original(client)
    checked = 0
    for finding in body["detailed_findings"]:
        clause_text = finding["clause_text"]
        start = finding["quote_start_offset"]
        end = finding["quote_end_offset"]
        if not start and not end:
            continue
        assert end > start, f"empty span for {finding['title']!r}"
        assert end <= len(clause_text), f"span runs past the clause for {finding['title']!r}"
        assert collapse(clause_text[start:end]) == collapse(finding["excerpt"])
        checked += 1
    assert checked >= 5, "most findings must carry a usable in-clause span"


def test_legacy_finding_shape_is_still_returned(client):
    register(client)
    body = analyse_original(client)
    assert len(body["findings"]) == len(body["detailed_findings"])
    for legacy, detailed in zip(body["findings"], body["detailed_findings"], strict=True):
        assert legacy["category"] == detailed["category"]
        assert legacy["severity"] == detailed["severity"]
        assert legacy["excerpt"] == detailed["excerpt"]
        assert legacy["explanation"] == detailed["explanation"]


# --------------------------------------------------------------------------- persistent statuses


def test_finding_status_persists_across_requests(client):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]

    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    resolved = client.patch(
        f"/api/review/findings/{finding_id}",
        json={"status": "resolved", "note": "Renegotiated on 12 March", "evidence": "Signed side letter"},
        headers={"X-CSRF-Token": csrf},
    )
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["status"] == "resolved"
    assert resolved.json()["status_note"] == "Renegotiated on 12 March"
    assert resolved.json()["status_changed_at"]

    reread = client.get(f"/api/review/findings/{finding_id}").json()
    assert reread["status"] == "resolved"
    assert reread["status_note"] == "Renegotiated on 12 March"

    detail = client.get(f"/api/review/{review['review_id']}").json()
    assert detail["counts"]["resolved"] >= 1
    assert detail["counts"]["open"] == len(detail["detailed_findings"]) - detail["counts"]["resolved"]


def test_resolving_a_finding_requires_a_note_or_evidence(client):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.patch(
        f"/api/review/findings/{finding_id}",
        json={"status": "resolved"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "evidence_required"


@pytest.mark.parametrize("status", ["open", "dismissed", "needs_professional_review"])
def test_supported_statuses_round_trip(client, status):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.patch(
        f"/api/review/findings/{finding_id}",
        json={"status": status, "note": "Reviewer decision"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == status
    assert client.get(f"/api/review/findings/{finding_id}").json()["status"] == status


def test_unknown_status_is_rejected(client):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.patch(
        f"/api/review/findings/{finding_id}",
        json={"status": "totally_fine"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 422


def test_status_update_requires_csrf(client):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    response = client.patch(
        f"/api/review/findings/{finding_id}",
        json={"status": "dismissed", "note": "no"},
    )
    assert response.status_code == 403


# --------------------------------------------------------------------------- version comparison


def test_revised_agreement_is_compared_against_the_original(client):
    register(client)
    document_id = make_document(client)
    first = analyse_original(client, document_id=str(document_id))
    second = analyse_revised(client, document_id, first["review_id"])

    assert second["parent_review_id"] == first["review_id"]
    assert second["round_number"] == first["round_number"] + 1

    comparison = client.get(f"/api/review/{second['review_id']}/comparison").json()
    assert comparison["previous_review_id"] == first["review_id"]
    assert comparison["previous_filename"] == ORIGINAL.name
    assert comparison["fixed_count"] == len(comparison["fixed"])
    assert comparison["unresolved_count"] == len(comparison["unresolved"])
    assert comparison["new_count"] == len(comparison["newly_introduced"])
    assert comparison["fixed_count"] + comparison["unresolved_count"] + comparison["new_count"] == len(
        comparison["items"]
    )
    assert "not legal advice" in comparison["disclaimer"]


def test_cured_concerns_are_reported_as_fixed_with_revised_evidence(client):
    register(client)
    document_id = make_document(client)
    first = analyse_original(client, document_id=str(document_id))
    second = analyse_revised(client, document_id, first["review_id"])

    comparison = client.get(f"/api/review/{second['review_id']}/comparison").json()
    fixed_titles = " ".join(item["title"] for item in comparison["fixed"])
    assert "90 days" in fixed_titles
    assert "non-refundable" in fixed_titles
    assert "carve-out" in fixed_titles
    for item in comparison["fixed"]:
        assert item["previous_finding_id"], "a fixed concern must link back to the earlier finding"
        assert item["evidence"].strip(), "a fixed concern must show evidence from the revision"
        assert "rewritten in the revised document" in item["evidence"]


def test_unchanged_clause_keeps_its_concern_open(client):
    register(client)
    document_id = make_document(client)
    first = analyse_original(client, document_id=str(document_id))
    second = analyse_revised(client, document_id, first["review_id"])

    comparison = client.get(f"/api/review/{second['review_id']}/comparison").json()
    standing = [item for item in comparison["unresolved"] if "Qualified legal review" in item["title"]]
    assert standing, "the standing professional-review limitation must stay open"
    assert "Still present" in standing[0]["summary"]
    assert "professional review" in standing[0]["evidence"].lower() or standing[0]["excerpt"]


def test_uploading_a_revision_stores_it_as_a_new_document_version(client):
    register(client)
    document_id = make_document(client)
    first = analyse_original(client, document_id=str(document_id))
    second = analyse_revised(client, document_id, first["review_id"])

    versions = client.get(f"/api/documents/{document_id}/versions").json()
    uploaded = [item for item in versions if item["change_note"].startswith("Analysed upload")]
    assert [item["version_number"] for item in uploaded] == [3, 2]
    assert uploaded[0]["change_note"] == f"Analysed upload: {REVISED.name}"
    assert first["version_number"] == 2
    assert second["version_number"] == 3
    assert second["document_version_id"] == uploaded[0]["id"]


def test_comparison_baseline_from_another_document_is_rejected(client):
    register(client)
    first_document = make_document(client, "First Agreement")
    other_document = make_document(client, "Second Agreement")
    first = analyse_original(client, document_id=str(first_document))
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": (REVISED.name, REVISED.read_bytes(), "text/plain")},
        data={"document_id": str(other_document), "previous_review_id": str(first["review_id"])},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "review_document_mismatch"


def test_reanalyze_reuses_the_latest_version_without_duplicating_it(client):
    register(client)
    document_id = make_document(client)
    first = analyse_original(client, document_id=str(document_id))
    versions_before = client.get(f"/api/documents/{document_id}/versions").json()

    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        f"/api/review/{first['review_id']}/reanalyze", headers={"X-CSRF-Token": csrf}
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["review"]["review_id"] != first["review_id"]
    assert body["review"]["parent_review_id"] == first["review_id"]
    assert body["review"]["version_number"] == 2
    assert body["review"]["document_version_id"] == first["document_version_id"]

    versions_after = client.get(f"/api/documents/{document_id}/versions").json()
    assert len(versions_after) == len(
        versions_before
    ), "re-analysing unchanged content must not create a duplicate version"


# --------------------------------------------------------------------------- quote location


@pytest.mark.parametrize(
    ("text", "quote", "expected"),
    [
        ("plain text", "plain text", "plain text"),
        ("one\ntwo   three", "one two three", "one\ntwo   three"),
        ("one\ntwo   three", "  one two  ", "one\ntwo"),
        ("first  second\n\nthird fourth", "second third", "second\n\nthird"),
        ("alpha\nbeta", "gamma", None),
        ("alpha beta", "alph", "alph"),
    ],
)
def test_quote_spans_survive_whitespace_drift(text, quote, expected):
    from app.services.clauses import find_quote_span

    span = find_quote_span(text, quote)
    if expected is None:
        assert span is None
        return
    assert span is not None
    assert text[span[0] : span[1]] == expected


def test_quote_offset_is_the_start_of_the_span():
    from app.services.clauses import find_quote_offset, find_quote_span

    text = "The Client shall pay\nwithin 30 days of the invoice."
    quote = "pay within 30 days"
    assert find_quote_offset(text, quote) == find_quote_span(text, quote)[0]


# --------------------------------------------------------------------------- per-finding Q&A


def test_finding_question_requires_a_configured_provider(client):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        f"/api/review/findings/{finding_id}/questions",
        json={"question": "Is the payment period enforceable?"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 503
    assert response.json()["detail"]["code"] == "ai_not_configured"


def test_finding_question_records_the_conversation(client, monkeypatch):
    register(client)
    monkeypatch.setattr("app.api.routes.review.get_ai_provider", AnsweringProvider)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]

    asked = client.post(
        f"/api/review/findings/{finding_id}/questions",
        json={"question": "Is the payment period enforceable?"},
        headers={"X-CSRF-Token": csrf},
    )
    assert asked.status_code == 200, asked.text
    body = asked.json()
    assert "Is the payment period enforceable?" in body["answer"]
    assert body["suggested_wording"] == review["detailed_findings"][0]["suggested_wording"]
    assert "not legal advice" in body["disclaimer"]
    assert [item["role"] for item in body["messages"]] == ["user", "assistant"]

    revision = client.post(
        f"/api/review/findings/{finding_id}/questions",
        json={"question": "What should it say instead?", "request_revision": True},
        headers={"X-CSRF-Token": csrf},
    )
    assert revision.status_code == 200, revision.text
    assert "within 30 days" in revision.json()["suggested_revision"]

    history = client.get(f"/api/review/findings/{finding_id}/questions").json()
    assert [item["role"] for item in history] == ["user", "assistant", "user", "assistant"]
    assert history[-1]["kind"] == "suggested_revision"
    assert "within 30 days" in history[-1]["suggested_revision"]


def test_question_history_is_scoped_to_the_question_thread(client, monkeypatch):
    register(client)
    monkeypatch.setattr("app.api.routes.review.get_ai_provider", AnsweringProvider)
    review = analyse_original(client)
    first_id = review["detailed_findings"][0]["id"]
    second_id = review["detailed_findings"][1]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    client.post(
        f"/api/review/findings/{first_id}/questions",
        json={"question": "About the first clause?"},
        headers={"X-CSRF-Token": csrf},
    )
    assert client.get(f"/api/review/findings/{second_id}/questions").json() == []


def test_empty_provider_answer_is_not_persisted(client, monkeypatch):
    register(client)
    monkeypatch.setattr("app.api.routes.review.get_ai_provider", SilentProvider)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        f"/api/review/findings/{finding_id}/questions",
        json={"question": "Anything?"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 502
    assert client.get(f"/api/review/findings/{finding_id}/questions").json() == []


# --------------------------------------------------------------------------- report and listing


def test_pdf_report_downloads_and_contains_the_review(client):
    register(client)
    review = analyse_original(client)
    response = client.get(f"/api/review/{review['review_id']}/report")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-type"] == "application/pdf"
    assert "review-report.pdf" in response.headers["content-disposition"]
    assert "attachment" in response.headers["content-disposition"]
    assert "no-store" in response.headers["cache-control"]
    assert len(response.content) > 2000


def test_pdf_report_covers_comparison_and_questions(client, monkeypatch):
    register(client)
    monkeypatch.setattr("app.api.routes.review.get_ai_provider", AnsweringProvider)
    document_id = make_document(client)
    first = analyse_original(client, document_id=str(document_id))
    second = analyse_revised(client, document_id, first["review_id"])

    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    client.post(
        f"/api/review/findings/{second['detailed_findings'][0]['id']}/questions",
        json={"question": "What is the risk here?"},
        headers={"X-CSRF-Token": csrf},
    )

    response = client.get(f"/api/review/{second['review_id']}/report")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_report_can_exclude_resolved_findings(client):
    register(client)
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    client.patch(
        f"/api/review/findings/{finding_id}",
        json={"status": "resolved", "note": "Agreed with counterparty"},
        headers={"X-CSRF-Token": csrf},
    )
    assert client.get(f"/api/review/{review['review_id']}/report?include_resolved=false").status_code == 200
    assert client.get(f"/api/review/{review['review_id']}/report?include_questions=false").status_code == 200


def test_review_list_and_detail_expose_counts(client):
    register(client)
    review = analyse_original(client)
    listing = client.get("/api/review")
    assert listing.status_code == 200
    body = listing.json()
    assert body["total"] == 1
    assert body["counts"]["total"] == len(review["detailed_findings"])
    assert body["items"][0]["id"] == review["review_id"]
    assert body["items"][0]["filename"] == ORIGINAL.name
    assert body["items"][0]["clause_count"] == review["clause_count"]

    detail = client.get(f"/api/review/{review['review_id']}").json()
    assert detail["clauses"], "clause references must survive a reload"
    assert detail["clauses"][0]["citation"]
    assert detail["comparison"] is not None
    assert detail["comparison"]["new_count"] == detail["counts"]["total"]


def test_list_filters_by_document(client):
    register(client)
    document_id = make_document(client)
    linked = analyse_original(client, document_id=str(document_id))
    standalone = analyse(client, ORIGINAL)
    filtered = client.get(f"/api/review?document_id={document_id}").json()
    assert filtered["total"] == 1
    assert filtered["items"][0]["id"] == linked["review_id"]
    assert client.get("/api/review").json()["total"] == 2
    assert standalone["document_id"] is None


def test_disclaimer_endpoint_is_not_captured_by_the_review_id_route(client):
    register(client)
    response = client.get("/api/review/disclaimer")
    assert response.status_code == 200
    assert "not legal advice" in response.json()["disclaimer"]


def test_deleting_a_review_removes_its_findings(client):
    register(client)
    review = analyse_original(client)
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    assert client.delete(f"/api/review/{review['review_id']}", headers={"X-CSRF-Token": csrf}).status_code == 204
    assert client.get(f"/api/review/{review['review_id']}").status_code == 404
    for finding in review["detailed_findings"]:
        assert client.get(f"/api/review/findings/{finding['id']}").status_code == 404


# --------------------------------------------------------------------------- privacy


def test_another_user_cannot_see_or_change_a_review(client):
    register(client, "owner@example.com")
    review = analyse_original(client)
    finding_id = review["detailed_findings"][0]["id"]

    client.post("/api/auth/logout")
    register(client, "intruder@example.com")

    assert client.get(f"/api/review/{review['review_id']}").status_code == 404
    assert client.get(f"/api/review/findings/{finding_id}").status_code == 404
    assert client.get(f"/api/review/{review['review_id']}/comparison").status_code == 404
    assert client.get(f"/api/review/{review['review_id']}/report").status_code == 404
    assert client.get(f"/api/review/findings/{finding_id}/questions").status_code == 404
    assert client.post(
        f"/api/review/findings/{finding_id}/questions",
        json={"question": "What does this say?"},
    ).status_code in {403, 404}
    assert client.get("/api/review").json()["total"] == 0
    assert client.get("/api/review?document_id=1").status_code == 200


def test_another_user_cannot_use_another_users_comparison_baseline(client):
    register(client, "owner@example.com")
    review = analyse_original(client)
    client.post("/api/auth/logout")
    register(client, "intruder@example.com")
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/review",
        files={"file": (REVISED.name, REVISED.read_bytes(), "text/plain")},
        data={"previous_review_id": str(review["review_id"])},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 404
    assert response.json()["detail"]["code"] == "review_not_found"


def test_unauthenticated_review_requests_are_rejected(client):
    assert client.get("/api/review").status_code == 401
    assert client.post("/api/review", files={"file": ("a.txt", b"x", "text/plain")}).status_code == 401
    assert client.get("/api/review/disclaimer").status_code == 200
