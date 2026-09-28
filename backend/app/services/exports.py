from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.services.disclaimers import REPORT_DISCLAIMER
from app.services.sanitizer import escape_text, html_to_text

SEVERITY_COLORS = {
    "high": colors.HexColor("#B42318"),
    "medium": colors.HexColor("#B54708"),
    "low": colors.HexColor("#175CD3"),
}
STATUS_LABELS = {
    "open": "Open",
    "resolved": "Resolved",
    "dismissed": "Dismissed",
    "needs_professional_review": "Needs professional review",
}
LIFECYCLE_LABELS = {
    "new": "New",
    "carried_over": "Carried over",
    "fixed": "Fixed in this version",
    "unresolved": "Unresolved",
    "newly_introduced": "Newly introduced",
}
CATEGORY_LABELS = {
    "unclear_clause": "Unclear clause",
    "missing_information": "Missing information",
    "inconsistent_detail": "Inconsistent detail",
    "ambiguous_wording": "Ambiguous wording",
    "unusual_obligation": "Unusual obligation",
    "professional_review": "Professional review",
}


def _blocks(content_html: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw in (
        content_html.replace("</p>", "</p>\n")
        .replace("</h2>", "</h2>\n")
        .replace("</h3>", "</h3>\n")
        .replace("</li>", "</li>\n")
        .splitlines()
    ):
        value = raw.strip()
        if not value:
            continue
        if value.startswith("<h2") or value.startswith("<h3"):
            tag = "h2" if value.startswith("<h2") else "h3"
            blocks.append((tag, html_to_text(value)))
        elif value.startswith("<li"):
            blocks.append(("li", html_to_text(value)))
        elif value:
            blocks.append(("p", html_to_text(value)))
    return blocks


def export_pdf(title: str, content_html: str) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.7 * inch,
        title=title,
        author="LegalEase",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="LegalTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=23,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#14213D"),
            spaceAfter=18,
        )
    )
    styles.add(
        ParagraphStyle(
            name="LegalH2",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#14213D"),
            spaceBefore=10,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="LegalH3",
            parent=styles["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13,
            textColor=colors.HexColor("#315EFB"),
            spaceBefore=7,
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            name="LegalBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            spaceAfter=7,
            textColor=colors.HexColor("#182230"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="LegalList",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            leftIndent=14,
            firstLineIndent=-9,
            spaceAfter=3,
            textColor=colors.HexColor("#182230"),
        )
    )
    story: list[object] = [Paragraph(escape_text(title), styles["LegalTitle"]), Spacer(1, 4)]
    for kind, text in _blocks(content_html):
        if not text:
            continue
        safe = escape_text(text).replace("\n", "<br/>")
        if kind == "h2":
            story.append(Paragraph(safe, styles["LegalH2"]))
        elif kind == "h3":
            story.append(Paragraph(safe, styles["LegalH3"]))
        elif kind == "li":
            story.append(Paragraph(f"• {safe}", styles["LegalList"]))
        else:
            story.append(Paragraph(safe, styles["LegalBody"]))
    if not story[2:]:
        story.append(Paragraph("This document is empty.", styles["LegalBody"]))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawCentredString(letter[0] / 2, 0.38 * inch, f"LegalEase  •  Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def export_docx(title: str, content_html: str) -> bytes:
    document = DocxDocument()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in _blocks(content_html):
        if not text:
            continue
        if kind == "h2":
            document.add_heading(text, level=1)
        elif kind == "h3":
            document.add_heading(text, level=2)
        elif kind == "li":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(text)
        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.add_run(text)
    document.core_properties.title = title
    document.core_properties.author = "LegalEase"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------- review report


def _review_styles() -> dict[str, ParagraphStyle]:
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=19,
            leading=24,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#14213D"),
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportSub",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#667085"),
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportH2",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12.5,
            leading=16,
            textColor=colors.HexColor("#14213D"),
            spaceBefore=13,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportH3",
            parent=styles["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13.5,
            textColor=colors.HexColor("#315EFB"),
            spaceBefore=9,
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            spaceAfter=6,
            textColor=colors.HexColor("#182230"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=12,
            spaceAfter=4,
            textColor=colors.HexColor("#475467"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportQuote",
            parent=styles["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=13,
            leftIndent=12,
            borderPadding=(4, 0, 4, 8),
            backColor=colors.HexColor("#F2F5FF"),
            textColor=colors.HexColor("#1D2939"),
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportDisclaimer",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=12.5,
            textColor=colors.HexColor("#7A2E0E"),
            spaceAfter=5,
        )
    )
    return styles


def _label(text: str) -> str:
    return f'<font color="#667085">{escape_text(text)}</font>'


def _page_reference(finding: Any) -> str:
    start = getattr(finding, "page_start", 1) or 1
    end = getattr(finding, "page_end", start) or start
    return f"Page {start}" if start == end else f"Pages {start}-{end}"


def _clause_reference(finding: Any) -> str:
    heading = (getattr(finding, "clause_heading", "") or "").strip()
    page = _page_reference(finding)
    return f"{heading} — {page}" if heading else page


def _finding_badge(finding: Any) -> str:
    severity = str(getattr(finding, "severity", "medium")).lower()
    color = SEVERITY_COLORS.get(severity, colors.HexColor("#175CD3"))
    severity_text = escape_text(severity.capitalize())
    status = escape_text(STATUS_LABELS.get(str(getattr(finding, "status", "open")), "Open"))
    lifecycle = LIFECYCLE_LABELS.get(str(getattr(finding, "lifecycle", "new")), "New")
    category = escape_text(
        CATEGORY_LABELS.get(str(getattr(finding, "category", "")), "Observation")
    )
    return (
        f'<font color="{color.hexval()}"><b>{severity_text} severity</b></font> · '
        f"{status} · {escape_text(lifecycle)} · {category}"
    )


def _highlighted_excerpt(finding: Any, clause_text: str) -> str:
    """Render the extracted clause with the flagged wording marked inside it."""

    excerpt = (getattr(finding, "excerpt", "") or "").strip()
    clause = (clause_text or "").strip()
    if not excerpt:
        return escape_text(clause[:600])
    start = getattr(finding, "quote_start_offset", 0) or 0
    end = getattr(finding, "quote_end_offset", 0) or 0
    if clause and 0 <= start < end <= len(clause) and clause[start:end].strip() == excerpt[: end - start]:
        prefix = escape_text(clause[:start])
        middle = f'<font color="{colors.HexColor("#B42318").hexval()}"><b>{escape_text(clause[start:end])}</b></font>'
        suffix = escape_text(clause[end : end + 700])
        return f"{prefix}{middle}{suffix}".strip()
    position = clause.find(excerpt[:200])
    if position < 0:
        return escape_text(excerpt)
    prefix = escape_text(clause[:position])
    middle = (
        f'<font color="{colors.HexColor("#B42318").hexval()}"><b>'
        f"{escape_text(excerpt)}</b></font>"
    )
    suffix = escape_text(clause[position + len(excerpt) : position + len(excerpt) + 700])
    return f"{prefix}{middle}{suffix}".strip()


def _clause_summary(clause: Any, limit: int = 180) -> str:
    """Read a clause summary from either a clause dataclass or an API schema object."""

    summary = getattr(clause, "summary", "")
    if callable(summary):
        summary = summary(limit)
    return escape_text(str(summary or "")[:limit])


def _clause_citation(clause: Any) -> str:
    """Prefer a ready-made citation when the clause carries one."""

    citation = str(getattr(clause, "citation", "") or "").strip()
    if citation:
        return escape_text(citation)
    page_start = int(getattr(clause, "page_start", 1) or 1)
    page_end = int(getattr(clause, "page_end", page_start) or page_start)
    return escape_text(
        f"Page {page_start}" if page_start == page_end else f"Pages {page_start}-{page_end}"
    )


def _finding_block(finding: Any, styles: dict[str, ParagraphStyle]) -> list[Any]:
    blocks: list[Any] = []
    title = escape_text(getattr(finding, "title", "") or "Untitled finding")
    blocks.append(Paragraph(f"{title}<br/>{_finding_badge(finding)}", styles["ReportH3"]))
    blocks.append(
        Paragraph(
            f"{_label('Source')}: {escape_text(_clause_reference(finding))}", styles["ReportSmall"]
        )
    )
    clause_text = getattr(finding, "clause_text", "") or ""
    if clause_text:
        blocks.append(
            Paragraph(
                f"{_label('Extracted clause')}: {_highlighted_excerpt(finding, clause_text)}",
                styles["ReportQuote"],
            )
        )
    explanation = (getattr(finding, "explanation", "") or "").strip()
    if explanation:
        blocks.append(Paragraph(f"{_label('Risk')}: {escape_text(explanation)}", styles["ReportBody"]))
    severity_explanation = (getattr(finding, "severity_explanation", "") or "").strip()
    if severity_explanation:
        blocks.append(
            Paragraph(
                f"{_label('Why this severity')}: {escape_text(severity_explanation)}",
                styles["ReportBody"],
            )
        )
    suggested = (getattr(finding, "suggested_wording", "") or "").strip()
    if suggested:
        blocks.append(
            Paragraph(f"{_label('Suggested wording')}: {escape_text(suggested)}", styles["ReportBody"])
        )
    question = (getattr(finding, "suggested_question", "") or "").strip()
    if question:
        blocks.append(
            Paragraph(
                f"{_label('Question for counsel')}: {escape_text(question)}", styles["ReportBody"]
            )
        )
    status_note = (getattr(finding, "status_note", "") or "").strip()
    if status_note:
        blocks.append(
            Paragraph(f"{_label('Reviewer note')}: {escape_text(status_note)}", styles["ReportSmall"])
        )
    status_evidence = (getattr(finding, "status_evidence", "") or "").strip()
    if status_evidence:
        blocks.append(
            Paragraph(
                f"{_label('Reviewer evidence')}: {escape_text(status_evidence)}",
                styles["ReportSmall"],
            )
        )
    resolution = (getattr(finding, "resolution_evidence", "") or "").strip()
    if resolution:
        blocks.append(
            Paragraph(
                f"{_label('Evidence that it was fixed')}: {escape_text(resolution)}",
                styles["ReportSmall"],
            )
        )
    return blocks


def export_review_report(
    *,
    review: Any,
    findings: list[Any],
    clauses: list[Any],
    comparison: Any = None,
    messages_by_finding: dict[int, list[Any]] | None = None,
    include_resolved: bool = True,
    include_questions: bool = True,
) -> bytes:
    """Build the exportable PDF review report for one clause-level review."""

    styles = _review_styles()
    messages_by_finding = messages_by_finding or {}
    selected = [
        finding
        for finding in findings
        if include_resolved or finding.status not in {"resolved", "dismissed"}
    ]
    selected.sort(
        key=lambda item: (
            {"high": 0, "medium": 1, "low": 2}.get(str(item.severity).lower(), 3),
            item.clause_index,
            item.id,
        )
    )

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=f"{review.document_title or review.filename} — LegalEase review report",
        author="LegalEase",
        subject="Automated document intelligence review report",
    )
    story: list[Any] = [
        Paragraph("LegalEase document review report", styles["ReportTitle"]),
        Paragraph(
            "Automated clause-level analysis · not legal advice", styles["ReportSub"]
        ),
    ]

    created = getattr(review, "created_at", None)
    story.append(Paragraph("1. Document details", styles["ReportH2"]))
    details = [
        ["Document", review.document_title or review.filename],
        ["Uploaded file", f"{review.filename} ({str(review.file_type).upper()})"],
        [
            "Version",
            f"Version {review.version_number}, review round {review.round_number}",
        ],
        [
            "Length",
            f"{review.page_count} page(s) · {review.word_count} words · "
            f"{review.clause_count} clause(s) extracted",
        ],
        [
            "Page references",
            "Source PDF page numbers"
            if review.page_reference_kind == "source_page"
            else "Estimated pages (the source format has no fixed pagination)",
        ],
        ["Analysed at", created.strftime("%Y-%m-%d %H:%M UTC") if created else "Not recorded"],
        ["Analysis engine", f"{review.provider} · {review.model}"],
    ]
    table = Table(details, colWidths=[1.5 * inch, 4.7 * inch], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#667085")),
                ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#E5EAF2")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    summary = (review.summary or "").strip()
    if summary:
        story.append(Spacer(1, 6))
        story.append(Paragraph(f"{_label('Summary')}: {escape_text(summary)}", styles["ReportBody"]))

    story.append(Paragraph("2. Extracted clauses", styles["ReportH2"]))
    if clauses:
        rows = [["Clause", "Source", "Extract"]]
        for clause in clauses:
            rows.append(
                [
                    escape_text(clause.heading or f"Clause {clause.index}"),
                    _clause_citation(clause),
                    _clause_summary(clause),
                ]
            )
        clause_table = Table(rows, colWidths=[1.7 * inch, 0.8 * inch, 3.7 * inch], repeatRows=1)
        clause_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F5FF")),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#E5EAF2")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(clause_table)
    else:
        story.append(
            Paragraph("No clause-level text was stored for this review.", styles["ReportBody"])
        )

    if comparison is not None and comparison.previous_review_id:
        story.append(Paragraph("3. Version comparison", styles["ReportH2"]))
        story.append(
            Paragraph(
                f'Compared against <b>{escape_text(comparison.previous_filename or "the earlier version")}</b>: '
                f"{comparison.fixed_count} fixed, {comparison.unresolved_count} unresolved, "
                f"{comparison.new_count} newly introduced.",
                styles["ReportBody"],
            )
        )
        if comparison.fixed:
            story.append(Paragraph("Fixed in this version", styles["ReportH3"]))
            story.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(
                                f"{escape_text(item.title)} — {escape_text(item.clause_citation)}. "
                                f"{escape_text(item.evidence or item.summary)}",
                                styles["ReportSmall"],
                            )
                        )
                        for item in comparison.fixed
                    ],
                    bulletType="bullet",
                    leftIndent=14,
                )
            )
        if comparison.unresolved:
            story.append(Paragraph("Still unresolved", styles["ReportH3"]))
            story.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(
                                f"{escape_text(item.title)} — {escape_text(item.clause_citation)}. "
                                f"{escape_text(item.summary)}",
                                styles["ReportSmall"],
                            )
                        )
                        for item in comparison.unresolved
                    ],
                    bulletType="bullet",
                    leftIndent=14,
                )
            )
        if comparison.newly_introduced:
            story.append(Paragraph("Newly introduced by this version", styles["ReportH3"]))
            story.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(
                                f"{escape_text(item.title)} — {escape_text(item.clause_citation)}. "
                                f"{escape_text(item.evidence or item.summary)}",
                                styles["ReportSmall"],
                            )
                        )
                        for item in comparison.newly_introduced
                    ],
                    bulletType="bullet",
                    leftIndent=14,
                )
            )
        findings_offset = 4
    else:
        findings_offset = 3

    story.append(
        Paragraph(f"{findings_offset}. Findings", styles["ReportH2"])
    )
    if selected:
        for index, finding in enumerate(selected, start=1):
            heading = f"{index}. {escape_text(finding.title or 'Finding')}"
            block = [Paragraph(heading, styles["ReportH3"])] + _finding_block(finding, styles)[1:]
            if include_questions:
                questions = messages_by_finding.get(finding.id, [])
                if questions:
                    block.append(Paragraph("Reviewer questions and answers", styles["ReportSmall"]))
                    for message in questions:
                        role = "Q" if message.role == "user" else "A"
                        block.append(
                            Paragraph(
                                f'<b>{role}</b> {escape_text(message.content)}',
                                styles["ReportSmall"],
                            )
                        )
                        if message.suggested_revision:
                            block.append(
                                Paragraph(
                                    f"{_label('Suggested revision')}: "
                                    f"{escape_text(message.suggested_revision)}",
                                    styles["ReportSmall"],
                                )
                            )
            story.append(KeepTogether(block))
    else:
        story.append(Paragraph("No findings were recorded for this review.", styles["ReportBody"]))

    story.append(Paragraph(f"{findings_offset + 1}. Unresolved questions", styles["ReportH2"]))
    questions = [
        item
        for finding in selected
        for item in [
            getattr(finding, "suggested_question", ""),
            getattr(finding, "status_note", ""),
        ]
        if str(item or "").strip()
    ]
    if getattr(review, "new_count", 0) and comparison is not None and comparison.new_count:
        questions.append(
            f"Which of the {comparison.new_count} newly introduced issue(s) are the result of the "
            "latest revision, and were they intended?"
        )
    if any(finding.status == "needs_professional_review" for finding in selected):
        questions.append(
            "Which items marked for professional review must a qualified lawyer resolve before signature?"
        )
    if not include_resolved:
        questions.append("Please confirm whether the excluded resolved and dismissed findings are closed.")
    if questions:
        story.append(
            ListFlowable(
                [ListItem(Paragraph(escape_text(str(item)), styles["ReportBody"])) for item in questions],
                bulletType="bullet",
                leftIndent=14,
            )
        )
    else:
        story.append(
            Paragraph("No outstanding questions were recorded.", styles["ReportBody"])
        )

    story.append(Paragraph(f"{findings_offset + 2}. Important notice", styles["ReportH2"]))
    story.append(
        Paragraph(
            f'<font backColor="#FFF6ED">{escape_text(REPORT_DISCLAIMER)}</font>',
            styles["ReportDisclaimer"],
        )
    )

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawCentredString(letter[0] / 2, 0.38 * inch, f"LegalEase  •  Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
