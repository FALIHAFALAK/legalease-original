from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.config import settings
from app.services.clauses import ExtractedDocument, paginate_text

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"
MAX_PDF_PAGES = 400
TEXT_LIMIT_MULTIPLIER = 0.8

ALLOWED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt"}
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "application/octet-stream",
    "text/plain",
    "text/markdown",
    "text/csv",
}
EXTENSION_CONTENT_TYPES = {
    "pdf": {"application/pdf", "application/octet-stream"},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
    "txt": {"text/plain", "text/markdown", "text/csv", "application/octet-stream"},
}


class FileValidationError(Exception):
    """Raised when an upload fails type, size, or content validation."""


def max_upload_bytes() -> int:
    return settings.max_upload_mb * 1024 * 1024


def normalize_extension(filename: str) -> str:
    return PurePath((filename or "").strip()).suffix.lower()


def validate_file_type(filename: str, content_type: str | None) -> str:
    """Return the validated file type, rejecting unsupported or mismatched uploads."""

    suffix = normalize_extension(filename)
    file_type = ALLOWED_EXTENSIONS.get(suffix)
    if file_type is None:
        raise FileValidationError("Only PDF, DOCX, and TXT files are supported.")
    declared = (content_type or "").split(";")[0].strip().lower()
    if declared:
        if declared not in ALLOWED_CONTENT_TYPES:
            raise FileValidationError("Only PDF, DOCX, and TXT files are supported.")
        if declared not in EXTENSION_CONTENT_TYPES[file_type]:
            raise FileValidationError(
                f"The file extension .{suffix.lstrip('.')} does not match the uploaded content type."
            )
    return file_type


def validate_file_size(data: bytes) -> None:
    if not data:
        raise FileValidationError("The uploaded file is empty.")
    limit = max_upload_bytes()
    if len(data) > limit:
        raise FileValidationError(
            f"Files must be smaller than {settings.max_upload_mb} MB. "
            f"This file is {round(len(data) / 1024 / 1024, 1)} MB."
        )


def _validate_pdf_signature(data: bytes) -> None:
    if not data.lstrip()[: len(PDF_MAGIC)] == PDF_MAGIC:
        raise FileValidationError("The file is not a readable PDF document.")


def _validate_docx_signature(data: bytes) -> None:
    if not data[: len(ZIP_MAGIC)] == ZIP_MAGIC:
        raise FileValidationError("The file is not a readable DOCX document.")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
    except zipfile.BadZipFile as exc:
        raise FileValidationError("The DOCX file could not be read safely.") from exc
    if "word/document.xml" not in names:
        raise FileValidationError("The DOCX file does not contain a readable document body.")


def _validate_text_payload(data: bytes) -> str:
    if b"\x00" in data[:8192]:
        raise FileValidationError("TXT files must be plain UTF-8 text, not binary data.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileValidationError("TXT files must use UTF-8 encoding.") from exc


def _extract_pdf(data: bytes) -> tuple[list[str], str]:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        pages = list(reader.pages)
    except Exception as exc:
        raise FileValidationError("The PDF could not be read safely.") from exc
    if not pages:
        raise FileValidationError("The PDF does not contain any pages.")
    if len(pages) > MAX_PDF_PAGES:
        raise FileValidationError(
            f"Documents must be {MAX_PDF_PAGES} pages or fewer for clause-level review."
        )
    extracted: list[str] = []
    for page in pages:
        try:
            extracted.append((page.extract_text() or "").strip())
        except Exception as exc:
            raise FileValidationError("The PDF text layer could not be read safely.") from exc
    if not any(extracted):
        raise FileValidationError(
            "No selectable text was found in the PDF. Scanned documents need OCR before review."
        )
    return extracted, "source_page"


def _extract_docx(data: bytes) -> str:
    try:
        document = DocxDocument(BytesIO(data))
    except Exception as exc:
        raise FileValidationError("The DOCX file could not be read safely.") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    blocks += [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    ]
    return "\n".join(blocks)


def extract_upload(
    filename: str, content_type: str | None, data: bytes
) -> ExtractedDocument:
    """Validate an upload and extract page-mapped text from PDF, DOCX, or TXT files."""

    file_type = validate_file_type(filename, content_type)
    validate_file_size(data)

    if file_type == "pdf":
        _validate_pdf_signature(data)
        pages, page_reference_kind = _extract_pdf(data)
        text = "\n\n".join(pages)
    elif file_type == "docx":
        _validate_docx_signature(data)
        text = _extract_docx(data).strip()
        pages = paginate_text(text) if text else [""]
        page_reference_kind = "estimated_page"
    else:
        text = _validate_text_payload(data).strip()
        pages = paginate_text(text) if text else [""]
        page_reference_kind = "estimated_page"

    text = text.strip()
    if not text:
        raise FileValidationError("No readable text was found in the uploaded file.")
    text_limit = max(1000, int(settings.gemini_max_input_chars * TEXT_LIMIT_MULTIPLIER))
    if len(text) > text_limit:
        raise FileValidationError(
            "The document is longer than the clause-review limit. "
            "Upload an extract that contains the clauses you want reviewed."
        )
    return ExtractedDocument(
        file_type=file_type, text=text, pages=pages, page_reference_kind=page_reference_kind
    )
